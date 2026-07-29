"""Excel 填写器 - 用 openpyxl 按模板填写周报 xlsx

模板布局见 templates.py
"""

from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from week_agent.config import DATA_DIR
from week_agent.weekly_report.templates import (
    TEMPLATE_LAYOUT,
    WEEKLY_TEMPLATE_PATH,
    validate_weekly_data,
)


def fill_weekly_report(
    reporter_name: str,
    week_range: str,
    this_week_work: list[dict],
    next_week_plan: list[dict],
    industry_info: list[dict],
    template_path: Optional[Path] = None,
    output_dir: Optional[Path] = None,
) -> Path:
    """按模板填写周报 xlsx

    Args:
        reporter_name: 填报人姓名（用于输出文件名）
        week_range: 周报区间，如 "2026.7.22-7.28"（用于文件名）
        this_week_work: 本周工作列表
        next_week_plan: 下周计划列表
        industry_info: 行业信息列表
        template_path: 模板路径（默认 WEEKLY_TEMPLATE_PATH）
        output_dir: 输出目录（默认 data/drafts/）

    Returns:
        生成的 xlsx 文件路径

    Raises:
        FileNotFoundError: 模板不存在
        ValueError: 数据校验失败
    """
    template = template_path or WEEKLY_TEMPLATE_PATH
    if not template.exists():
        raise FileNotFoundError(f"周报模板不存在: {template}")

    # 校验数据
    data = {
        "this_week_work": this_week_work,
        "next_week_plan": next_week_plan,
        "industry_info": industry_info,
    }
    errors = validate_weekly_data(data)
    if errors:
        raise ValueError("周报数据校验失败: " + "; ".join(errors))

    # 加载模板（保留样式与数据验证）
    wb = load_workbook(str(template))
    ws = wb.active

    # 1. 填写本周工作
    layout = TEMPLATE_LAYOUT["this_week_work"]
    rows = list(layout["rows"])
    for i, item in enumerate(this_week_work[: layout["max_rows"]]):
        row = rows[i]
        ws[f"{layout['columns']['category']}{row}"] = item.get("category", "")
        ws[f"{layout['columns']['task']}{row}"] = item.get("task", "")
        ws[f"{layout['columns']['status']}{row}"] = item.get("status", "")
        ws[f"{layout['columns']['remark']}{row}"] = item.get("remark", "")

    # 2. 填写下周计划
    layout = TEMPLATE_LAYOUT["next_week_plan"]
    rows = list(layout["rows"])
    for i, item in enumerate(next_week_plan[: layout["max_rows"]]):
        row = rows[i]
        ws[f"{layout['columns']['category']}{row}"] = item.get("category", "")
        ws[f"{layout['columns']['task']}{row}"] = item.get("task", "")
        ws[f"{layout['columns']['status']}{row}"] = item.get("status", "")
        ws[f"{layout['columns']['risk']}{row}"] = item.get("risk", "")

    # 3. 填写行业信息
    layout = TEMPLATE_LAYOUT["industry_info"]
    rows = list(layout["rows"])
    for i, item in enumerate(industry_info[: layout["max_rows"]]):
        row = rows[i]
        ws[f"{layout['columns']['content']}{row}"] = item.get("content", "")
        ws[f"{layout['columns']['remark']}{row}"] = item.get("remark", "")

    # 保存
    out_dir = output_dir or (DATA_DIR / "drafts")
    out_dir.mkdir(parents=True, exist_ok=True)
    # 文件名去除可能非法的字符
    safe_name = "".join(c if c.isalnum() or c in "._-" else "_" for c in reporter_name)
    # 周报区间只取最后一天作为文件名日期：
    # "2026.7.27-7.28" / "2026.7.27~2026.7.28" / "2026.7.27 到 2026.7.28" → "2026.7.28"
    safe_range = _extract_last_day(week_range)
    filename = f"工作周报-{safe_name}-{safe_range}.xlsx"
    output_path = out_dir / filename
    wb.save(str(output_path))
    return output_path


def _extract_last_day(week_range: str) -> str:
    """从周报区间提取最后一天的日期，格式为 YYYY.M.D

    示例：
        "2026.7.27-7.28"      → "2026.7.28"
        "2026.7.27~2026.7.28" → "2026.7.28"
        "2026.7.27 到 2026.7.28" → "2026.7.28"
        "2026.7.27"           → "2026.7.27"（单日期原样返回）
    """
    import re

    if not week_range:
        return "unknown"

    s = week_range.strip()

    # 匹配 "起始日期 分隔符 结束日期"，分隔符为 - ~ 到 至 等
    m = re.search(
        r"(\d{4}\.\d{1,2}\.\d{1,2})\s*[-~到至]\s*(\d{1,2}\.\d{1,2}|\d{4}\.\d{1,2}\.\d{1,2})",
        s,
    )
    if m:
        start = m.group(1)
        end = m.group(2)
        # 结尾无年份时补上开头年份
        if len(end.split(".")) == 2:
            year = start.split(".")[0]
            end = f"{year}.{end}"
        return end

    # 没匹配到范围，尝试匹配单个日期
    m = re.search(r"(\d{4}\.\d{1,2}\.\d{1,2})", s)
    if m:
        return m.group(1)

    # 兜底：清理非法字符后原样返回
    return "".join(c if c.isalnum() or c in "._-" else "_" for c in s)
