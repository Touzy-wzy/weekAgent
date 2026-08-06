"""读取上传文件工具（通用智能体对话用）

支持白名单类型：.docx / .pdf / .txt / .md / .xlsx
与周报专用 ReadUploadedDocTool（仅 .docx）并存：
- read_uploaded_doc：周报会话专用，参数 session_id
- read_uploaded_file：通用对话专用，参数 session_id + filename
"""

from pathlib import Path
from typing import Any, Dict, List

from hello_agents.tools.base import Tool, ToolParameter
from hello_agents.tools.errors import ToolErrorCode
from hello_agents.tools.response import ToolResponse

from week_agent.config import DATA_DIR
from week_agent.weekly_report.doc_parser import parse_docx_to_text

# 允许上传/读取的文件扩展名白名单
ALLOWED_EXTENSIONS = {".docx", ".pdf", ".txt", ".md", ".xlsx"}

# 返回给 LLM 的最大字符数（避免上下文爆炸）
MAX_CONTENT_CHARS = 12000


def parse_uploaded_file(file_path: Path) -> str:
    """按扩展名解析文件为纯文本

    Args:
        file_path: 上传文件路径

    Returns:
        纯文本内容

    Raises:
        ValueError: 不支持的格式或解析失败
    """
    file_path = Path(file_path)
    ext = file_path.suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"不支持的文件类型: {ext}（允许: {', '.join(sorted(ALLOWED_EXTENSIONS))}）")

    if ext == ".docx":
        return parse_docx_to_text(file_path)

    if ext in (".txt", ".md"):
        return file_path.read_text(encoding="utf-8", errors="replace")

    if ext == ".xlsx":
        return _parse_xlsx_to_text(file_path)

    if ext == ".pdf":
        return _parse_pdf_to_text(file_path)

    raise ValueError(f"不支持的文件类型: {ext}")


def _parse_xlsx_to_text(file_path: Path) -> str:
    """用 openpyxl 解析 .xlsx 为纯文本（工作表 + 单元格内容）"""
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"【工作表：{ws.title}】")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    if not parts:
        raise ValueError("Excel 文件内容为空")
    return "\n".join(parts)


def _parse_pdf_to_text(file_path: Path) -> str:
    """用 pypdf 解析 .pdf 为纯文本"""
    from pypdf import PdfReader

    reader = PdfReader(str(file_path))
    parts: list[str] = []
    for i, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text.strip():
            parts.append(f"【第 {i} 页】\n{text.strip()}")
    if not parts:
        raise ValueError("PDF 内容为空（可能是扫描件，不支持 OCR）")
    return "\n\n".join(parts)


class ReadUploadedFileTool(Tool):
    """读取通用对话会话中用户上传的文件，返回纯文本"""

    def __init__(self):
        super().__init__(
            name="read_uploaded_file",
            description=(
                "读取指定会话中用户上传的文件内容（支持 .docx/.pdf/.txt/.md/.xlsx），"
                "返回纯文本。用于用户上传资料后，智能体阅读并分析这些资料。"
                "先调用 list_uploaded_files 查看会话已有哪些文件。"
            ),
        )

    def get_parameters(self) -> List[ToolParameter]:
        return [
            ToolParameter(
                name="session_id",
                type="string",
                description="智能体对话会话 ID（上传文件时绑定的 session_id）",
                required=True,
            ),
            ToolParameter(
                name="filename",
                type="string",
                description="要读取的文件名（含扩展名），如 report.docx",
                required=True,
            ),
        ]

    def run(self, parameters: Dict[str, Any]) -> ToolResponse:
        session_id = str(parameters.get("session_id", "")).strip()
        filename = str(parameters.get("filename", "")).strip()
        if not session_id:
            return ToolResponse.error(
                code=ToolErrorCode.INVALID_PARAM,
                message="session_id 不能为空",
            )
        if not filename:
            return ToolResponse.error(
                code=ToolErrorCode.INVALID_PARAM,
                message="filename 不能为空，先调用 list_uploaded_files 查看可用文件",
            )

        upload_dir = DATA_DIR / "uploads" / session_id
        if not upload_dir.exists():
            return ToolResponse.error(
                code=ToolErrorCode.NOT_FOUND,
                message=f"会话 {session_id} 未上传任何文件",
            )

        file_path = upload_dir / Path(filename).name  # 防路径穿越
        if not file_path.exists() or not file_path.is_file():
            return ToolResponse.error(
                code=ToolErrorCode.NOT_FOUND,
                message=f"文件 {filename} 不存在，可用文件: {', '.join(p.name for p in upload_dir.iterdir() if p.is_file())}",
            )

        print(f"📄 [read_uploaded_file] 读取: {file_path.name}")
        try:
            text = parse_uploaded_file(file_path)
            if not text.strip():
                return ToolResponse.partial(
                    text=f"文件 {file_path.name} 内容为空",
                    data={"session_id": session_id, "file": file_path.name, "text": ""},
                )
            preview = text[:MAX_CONTENT_CHARS]
            truncated = len(text) > MAX_CONTENT_CHARS
            print(f"  ✅ 解析成功，{len(text)} 字符" + ("（已截断）" if truncated else ""))
            return ToolResponse.success(
                text=f"文件 {file_path.name} 解析成功（{len(text)} 字符）"
                     + ("，已截断前 " + str(MAX_CONTENT_CHARS) + " 字符" if truncated else "") + "：\n\n"
                     + preview,
                data={
                    "session_id": session_id,
                    "file": file_path.name,
                    "text": preview,
                    "length": len(text),
                    "truncated": truncated,
                },
            )
        except Exception as e:
            print(f"  ❌ 解析失败: {e}")
            return ToolResponse.error(
                code=ToolErrorCode.EXECUTION_ERROR,
                message=f"文件解析失败: {str(e)}",
                context={"session_id": session_id, "file": file_path.name},
            )


class ListUploadedFilesTool(Tool):
    """列出指定会话已上传的文件"""

    def __init__(self):
        super().__init__(
            name="list_uploaded_files",
            description=(
                "列出指定会话中用户已上传的所有文件（文件名 + 大小 + 类型）。"
                "当用户提到上传了文件、或询问当前会话有哪些文件时调用。"
            ),
        )

    def get_parameters(self) -> List[ToolParameter]:
        return [
            ToolParameter(
                name="session_id",
                type="string",
                description="智能体对话会话 ID",
                required=True,
            ),
        ]

    def run(self, parameters: Dict[str, Any]) -> ToolResponse:
        session_id = str(parameters.get("session_id", "")).strip()
        if not session_id:
            return ToolResponse.error(
                code=ToolErrorCode.INVALID_PARAM,
                message="session_id 不能为空",
            )

        upload_dir = DATA_DIR / "uploads" / session_id
        if not upload_dir.exists():
            return ToolResponse.success(
                text="会话未上传任何文件",
                data={"session_id": session_id, "files": []},
            )

        files = []
        for p in upload_dir.iterdir():
            if p.is_file():
                files.append({
                    "name": p.name,
                    "size": p.stat().st_size,
                    "ext": p.suffix.lower(),
                })
        files.sort(key=lambda f: f["name"])
        if not files:
            return ToolResponse.success(
                text="会话未上传任何文件",
                data={"session_id": session_id, "files": []},
            )

        lines = [f"- {f['name']}（{f['size']} 字节）" for f in files]
        return ToolResponse.success(
            text=f"会话已上传 {len(files)} 个文件：\n" + "\n".join(lines),
            data={"session_id": session_id, "files": files},
        )
