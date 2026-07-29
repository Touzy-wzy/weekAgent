"""周报智能体验收测试

覆盖：
1. 模板 schema 校验
2. Excel 填写器
3. Word 解析器
4. 会话状态机
5. Web API（TestClient）
"""

import json
import sys
import traceback
from pathlib import Path

# 项目根目录加入 sys.path
ROOT = Path(__file__).parent.resolve()
sys.path.insert(0, str(ROOT))

from week_agent.weekly_report.templates import (
    CATEGORY_ENUM,
    NEXT_WEEK_STATUS_ENUM,
    THIS_WEEK_STATUS_ENUM,
    WEEKLY_TEMPLATE_PATH,
    validate_weekly_data,
)
from week_agent.weekly_report.excel_filler import fill_weekly_report
from week_agent.weekly_report.doc_parser import parse_docx_to_text
from week_agent.weekly_report.session import SessionState, WeeklySession
from week_agent.weekly_report.agent import create_weekly_tool_registry

results: list[tuple[str, bool, str]] = []


def case(name):
    def deco(fn):
        def wrapper():
            try:
                fn()
                results.append((name, True, "OK"))
                print(f"[PASS] {name}")
            except Exception as e:
                tb = traceback.format_exc()
                results.append((name, False, str(e) + "\n" + tb))
                print(f"[FAIL] {name}: {e}")
        return wrapper
    return deco


# ---------- 1. 模板校验 ----------

@case("模板文件存在")
def test_template_exists():
    assert WEEKLY_TEMPLATE_PATH.exists(), f"模板不存在: {WEEKLY_TEMPLATE_PATH}"


@case("枚举值与模板下拉一致")
def test_enum_values():
    assert CATEGORY_ENUM == ["61850", "AI", "其他"]
    assert THIS_WEEK_STATUS_ENUM == ["已完成", "进行中", "已延期"]
    assert NEXT_WEEK_STATUS_ENUM == ["完成", "开展", "延期"]


@case("schema 校验通过合法数据")
def test_validate_valid():
    data = {
        "this_week_work": [
            {"category": "AI", "task": "学习智能体", "status": "进行中", "remark": ""}
        ],
        "next_week_plan": [
            {"category": "61850", "task": "开发", "status": "开展", "risk": ""}
        ],
        "industry_info": [
            {"content": "OpenAI 发布新模型", "remark": ""}
        ],
    }
    errors = validate_weekly_data(data)
    assert errors == [], f"不应有错误: {errors}"


@case("schema 校验拒绝非法枚举")
def test_validate_invalid_enum():
    data = {
        "this_week_work": [
            {"category": "InvalidCat", "task": "x", "status": "InvalidStatus"}
        ],
        "next_week_plan": [],
        "industry_info": [],
    }
    errors = validate_weekly_data(data)
    assert len(errors) >= 2, f"应至少 2 个错误: {errors}"


@case("schema 校验拒绝超长数组")
def test_validate_too_many():
    data = {
        "this_week_work": [{"category": "AI", "task": "x", "status": "进行中"}] * 6,
        "next_week_plan": [],
        "industry_info": [],
    }
    errors = validate_weekly_data(data)
    assert any("超过最大行数" in e for e in errors), f"应报超长: {errors}"


# ---------- 2. Excel 填写 ----------

@case("Excel 填写并验证内容")
def test_fill_excel():
    this_week = [
        {"category": "AI", "task": "学习 hello-agents 框架", "status": "已完成", "remark": ""},
        {"category": "61850", "task": "积成电子项目周报系统设计", "status": "进行中", "remark": "需协调前端"},
        {"category": "其他", "task": "新员工培训", "status": "已延期", "remark": "延期至下周"},
    ]
    next_week = [
        {"category": "AI", "task": "智能体案例调研", "status": "开展", "risk": "低"},
        {"category": "61850", "task": "周报系统开发", "status": "完成", "risk": ""},
    ]
    industry = [
        {"content": "OpenAI 发布 GPT-5", "remark": "关注多模态能力"},
        {"content": "Anthropic 推出 MCP 协议", "remark": ""},
    ]
    out = fill_weekly_report(
        reporter_name="测试用户",
        week_range="2026.7.22-7.28",
        this_week_work=this_week,
        next_week_plan=next_week,
        industry_info=industry,
    )
    assert out.exists(), f"输出文件不存在: {out}"
    assert out.name == "工作周报-测试用户-2026.7.22-7.28.xlsx"

    # 读回校验
    from openpyxl import load_workbook
    wb = load_workbook(str(out), data_only=True)
    ws = wb.active
    # 本周工作第一行 (row 4)
    assert ws["B4"].value == "AI", f"B4 应为 AI, 实际 {ws['B4'].value}"
    assert ws["C4"].value == "学习 hello-agents 框架"
    assert ws["D4"].value == "已完成"
    # 第三行 (row 6)
    assert ws["B6"].value == "其他"
    assert ws["D6"].value == "已延期"
    # 下周计划第一行 (row 12)
    assert ws["B12"].value == "AI"
    assert ws["D12"].value == "开展"
    # 行业信息第一行 (row 21)
    assert ws["B21"].value == "OpenAI 发布 GPT-5"
    assert ws["E21"].value == "关注多模态能力"
    print(f"  生成文件: {out}")


@case("Excel 填写拒绝非法数据")
def test_fill_excel_invalid():
    try:
        fill_weekly_report(
            reporter_name="x",
            week_range="x",
            this_week_work=[{"category": "BAD", "task": "x", "status": "进行中"}],
            next_week_plan=[],
            industry_info=[],
        )
        assert False, "应抛 ValueError"
    except ValueError as e:
        assert "校验失败" in str(e)


# ---------- 3. Word 解析 ----------

@case("Word 解析器生成并读取测试文件")
def test_docx_parser():
    from docx import Document
    # 临时生成一个 docx
    tmp = ROOT / "data" / "test_upload.docx"
    tmp.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("本周工作总结", level=1)
    doc.add_paragraph("1. 完成了 hello-agents 框架的集成")
    doc.add_paragraph("2. 设计了周报智能体方案")
    table = doc.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "状态"
    table.rows[0].cells[2].text = "负责人"
    table.rows[1].cells[0].text = "周报系统"
    table.rows[1].cells[1].text = "进行中"
    table.rows[1].cells[2].text = "王兆阳"
    doc.save(str(tmp))

    text = parse_docx_to_text(tmp)
    assert "本周工作总结" in text
    assert "hello-agents 框架的集成" in text
    assert "周报系统 | 进行中 | 王兆阳" in text
    tmp.unlink()


@case("Word 解析器报错：文件不存在")
def test_docx_not_found():
    try:
        parse_docx_to_text(Path("nonexistent.docx"))
        assert False, "应抛 FileNotFoundError"
    except FileNotFoundError:
        pass


# ---------- 4. 会话状态机 ----------

@case("会话创建与持久化")
def test_session_persist():
    s = WeeklySession(reporter_name="张三", week_range="2026.7.22-7.28")
    s.save()
    loaded = WeeklySession.load(s.session_id)
    assert loaded is not None
    assert loaded.reporter_name == "张三"
    assert loaded.state == SessionState.INIT
    # 清理
    s.delete()
    assert WeeklySession.load(s.session_id) is None


@case("会话状态流转")
def test_session_transition():
    s = WeeklySession()
    assert s.state == SessionState.INIT
    s.transition(SessionState.GATHERING)
    assert s.state == SessionState.GATHERING
    s.transition(SessionState.DRAFTING)
    s.transition(SessionState.FILLING)
    s.transition(SessionState.REVIEWING)
    s.transition(SessionState.COLLECTING_RECIPIENT)
    s.transition(SessionState.CONFIRMING_SEND)
    s.transition(SessionState.SENT)
    assert s.state == SessionState.SENT
    s.delete()


@case("会话 INIT 缺失字段检查")
def test_session_missing():
    s = WeeklySession()
    missing = s.check_init_complete()
    assert "reporter_name" in missing
    assert "week_range" in missing
    assert "data_source" in missing
    # 部分填充
    s.reporter_name = "李四"
    s.data_source = "flowus"
    missing = s.check_init_complete()
    assert "reporter_name" not in missing
    assert "flowus_project" in missing  # 选了 flowus 但没项目名
    s.delete()


# ---------- 5. 工具注册 ----------

@case("周报 Agent 工具集注册完整")
def test_tool_registry():
    reg = create_weekly_tool_registry()
    # ToolRegistry 提供 list_tools() 返回工具名列表
    if hasattr(reg, "list_tools"):
        names = list(reg.list_tools())
    elif hasattr(reg, "get_all_tools"):
        names = [t.name for t in reg.get_all_tools()]
    else:
        names = []
    expected = [
        "flowus_search", "flowus_list_pages", "flowus_get_page",
        "flowus_fetch_weekly", "read_uploaded_doc", "fill_weekly_excel",
        "agently_compose_mail", "agently_send_mail",
    ]
    for n in expected:
        assert n in names, f"缺少工具: {n}（实际 {names}）"


# ---------- 6. Web API ----------

@case("Web API：周报会话 CRUD")
def test_web_api_session():
    from fastapi.testclient import TestClient
    from week_agent.web.app import app
    client = TestClient(app)

    # 创建
    r = client.post("/api/weekly/sessions")
    assert r.status_code == 200
    sid = r.json()["session_id"]
    assert r.json()["state"] == "init"

    # 查询
    r = client.get(f"/api/weekly/sessions/{sid}")
    assert r.status_code == 200
    assert r.json()["session_id"] == sid

    # 列表
    r = client.get("/api/weekly/sessions")
    assert r.status_code == 200
    assert any(s["session_id"] == sid for s in r.json())

    # 清理
    s = WeeklySession.load(sid)
    if s:
        s.delete()


@case("Web API：收件人管理")
def test_web_api_recipients():
    from fastapi.testclient import TestClient
    from week_agent.web.app import app
    from week_agent.web.weekly_services import RECIPIENTS_FILE
    # 备份原文件
    backup = RECIPIENTS_FILE.read_bytes() if RECIPIENTS_FILE.exists() else None
    try:
        client = TestClient(app)
        # 添加
        r = client.post("/api/weekly/recipients", params={"name": "测试张三", "email": "test_zs@example.com"})
        assert r.status_code == 200
        emails = [x["email"] for x in r.json()]
        assert "test_zs@example.com" in emails
        # 列表
        r = client.get("/api/weekly/recipients")
        assert r.status_code == 200
        assert any(x["email"] == "test_zs@example.com" for x in r.json())
        # 删除
        r = client.delete("/api/weekly/recipients", params={"email": "test_zs@example.com"})
        assert r.status_code == 200
        emails = [x["email"] for x in r.json()]
        assert "test_zs@example.com" not in emails
    finally:
        # 还原
        if backup is not None:
            RECIPIENTS_FILE.write_bytes(backup)
        elif RECIPIENTS_FILE.exists():
            RECIPIENTS_FILE.unlink()


@case("Web API：上传 Word 文件")
def test_web_api_upload():
    from fastapi.testclient import TestClient
    from io import BytesIO
    from docx import Document
    from week_agent.web.app import app
    from week_agent.config import DATA_DIR

    client = TestClient(app)
    # 创建会话
    r = client.post("/api/weekly/sessions")
    sid = r.json()["session_id"]
    # 生成测试 docx
    doc = Document()
    doc.add_paragraph("测试内容")
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)

    # 上传
    r = client.post(
        f"/api/weekly/sessions/{sid}/upload",
        files={"file": ("test.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert r.status_code == 200, r.text
    assert r.json()["filename"] == "test.docx"
    # 文件应存在
    f = DATA_DIR / "uploads" / sid / "test.docx"
    assert f.exists()
    # 会话应已绑定 data_source=docx
    r = client.get(f"/api/weekly/sessions/{sid}")
    assert r.json()["data_source"] == "docx"

    # 清理
    s = WeeklySession.load(sid)
    if s:
        s.delete()


@case("Web API：拒绝非 docx 上传")
def test_web_api_reject_non_docx():
    from fastapi.testclient import TestClient
    from io import BytesIO
    from week_agent.web.app import app

    client = TestClient(app)
    r = client.post("/api/weekly/sessions")
    sid = r.json()["session_id"]
    buf = BytesIO(b"hello")
    r = client.post(
        f"/api/weekly/sessions/{sid}/upload",
        files={"file": ("bad.txt", buf, "text/plain")},
    )
    assert r.status_code == 400
    s = WeeklySession.load(sid)
    if s:
        s.delete()


@case("Web API：approve 状态校验")
def test_web_api_approve_guard():
    from fastapi.testclient import TestClient
    from week_agent.web.app import app

    client = TestClient(app)
    r = client.post("/api/weekly/sessions")
    sid = r.json()["session_id"]
    # 初始 INIT 状态 approve 应失败
    r = client.post(f"/api/weekly/sessions/{sid}/approve")
    assert r.status_code == 400
    s = WeeklySession.load(sid)
    if s:
        s.delete()


# ---------- 主函数 ----------

def main():
    print("\n" + "=" * 60)
    print("周报智能体验收测试")
    print("=" * 60)

    tests = [
        test_template_exists,
        test_enum_values,
        test_validate_valid,
        test_validate_invalid_enum,
        test_validate_too_many,
        test_fill_excel,
        test_fill_excel_invalid,
        test_docx_parser,
        test_docx_not_found,
        test_session_persist,
        test_session_transition,
        test_session_missing,
        test_tool_registry,
        test_web_api_session,
        test_web_api_recipients,
        test_web_api_upload,
        test_web_api_reject_non_docx,
        test_web_api_approve_guard,
    ]
    for t in tests:
        t()

    print("\n" + "=" * 60)
    passed = sum(1 for _, ok, _ in results if ok)
    failed = sum(1 for _, ok, _ in results if not ok)
    print(f"结果: {passed} 通过 / {failed} 失败 / 共 {len(results)}")
    if failed:
        print("\n失败详情:")
        for name, ok, msg in results:
            if not ok:
                print(f"\n[FAIL] {name}\n{msg}\n")
        sys.exit(1)
    print("✅ 全部通过")


if __name__ == "__main__":
    main()
